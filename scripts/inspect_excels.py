import os
import openpyxl

paths = [
    r"c:\Users\sakth\OneDrive\Documents\Arch realty\IT\company project\All Format ( IT Work )",
    r"c:\Users\sakth\OneDrive\Documents\Arch realty\IT\company project\Excel data for software"
]

for base_path in paths:
    print("="*60)
    print("DIRECTORY:", base_path)
    print("="*60)
    if not os.path.exists(base_path):
        print("Path does not exist!")
        continue
    for f in os.listdir(base_path):
        if f.endswith('.xlsx') or f.endswith('.xls'):
            full_path = os.path.join(base_path, f)
            try:
                wb = openpyxl.load_workbook(full_path, read_only=True, data_only=True)
                sheetnames = wb.sheetnames
                print(f"\nFILE: {f}")
                print(f"  Sheets: {sheetnames}")
                for sname in sheetnames[:3]:
                    sheet = wb[sname]
                    # Read first 3 rows
                    rows = []
                    for row in sheet.iter_rows(max_row=3, values_only=True):
                        if any(row):
                            rows.append([str(c) if c is not None else '' for c in row[:15]])
                    if rows:
                        print(f"    Sheet '{sname}' Row 1: {rows[0]}")
                        if len(rows) > 1:
                            print(f"    Sheet '{sname}' Row 2: {rows[1]}")
                wb.close()
            except Exception as e:
                print(f"FILE {f} error: {e}")
