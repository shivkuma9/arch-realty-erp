import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

p = r"c:\Users\sakth\OneDrive\Documents\Arch realty\IT\company project\Excel data for software\ARCH_REALTY_BDE.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
for sname in wb.sheetnames:
    sheet = wb[sname]
    print(f"\n--- Sheet: {sname} (Rows: {sheet.max_row}) ---")
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i > 4:
            break
        vals = [str(c).strip() if c is not None else '' for c in row if c is not None]
        print(f"Row {i+1}: {vals[:10]}")
wb.close()
